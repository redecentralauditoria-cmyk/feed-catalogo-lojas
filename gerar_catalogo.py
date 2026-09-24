"""
Gera o catalogo.csv do catalogo da Meta a partir das planilhas do FarmaPRO.

REGRAS DE SEGURANCA (nao alterar sem falar com o Marcel):
- So entram produtos cujo "Nome Grupo" comeca com HPC (higiene pessoal e cosmeticos).
  Os grupos LB / GENER / SIMIL sao MEDICAMENTO pelo cadastro oficial e NUNCA entram.
  Esse filtro e a trava real: o medicamento nao existe no arquivo, entao a IA nao
  tem como informar preco dele pelo catalogo.
- Preco = coluna P.M.C. do CADASTRO (ja com desconto Fidelidade).
"""
import openpyxl, csv, glob, os, re, sys, unicodedata
from datetime import datetime

BASE = r"C:\Users\marcel.pereira\Desktop\AUDITORIAS\01-CLAUDE"
# Cadastro e Estoque vem da pasta dedicada deste projeto (o Marcel cola os arquivos aqui manualmente,
# nunca direto da rede) - pasta com espacos duplos de proposito, e o nome real da pasta no disco.
AUTOMATIZ_DIR = r"C:\Users\marcel.pereira\CLAUDE\AUTOMATIZ. ATENDIM.  WHATSAPP  LOJAS"
CADASTRO_DIR = AUTOMATIZ_DIR
ESTOQUE_DIR = AUTOMATIZ_DIR
BANCO_FOTOS = os.path.join(BASE, "IMAGENS-CATALOGO", "BANCO_IMAGENS_INSTABUY.xlsx")
SAIDA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "catalogo.csv")
LINK_LOJA = "https://loja.redecentralfarma.com.br/"

EXCLUIR_NOME = ["violeta genciana"]  # classificado como "para infeccoes" na loja online


def mais_recente(pasta, padrao):
    achados = glob.glob(os.path.join(pasta, padrao))
    if not achados:
        sys.exit(f"ERRO: nenhum arquivo '{padrao}' em {pasta}")
    return max(achados, key=os.path.getmtime)


def norm(s):
    s = unicodedata.normalize("NFD", str(s))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip().lower()


print("Gerando catalogo...\n")

# ---------- 1. fotos publicas (banco Instabuy) ----------
wb = openpyxl.load_workbook(BANCO_FOTOS, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
fotos = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    cod, nome, marca, img = r[1], r[4], r[5], r[7]
    if img and img != "Sem imagem" and cod:
        fotos[str(cod).strip()] = (nome, marca, img)
print(f"  fotos publicas disponiveis: {len(fotos)}")

# ---------- 2. estoque (so entra o que existe em alguma loja) ----------
arq_est = mais_recente(ESTOQUE_DIR, "Estoque das Filiais em *.XLSX")
wb = openpyxl.load_workbook(arq_est, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
tem_estoque = set()
for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True)):
    cod = r[0]
    try:
        total = float(r[3] or 0)
    except (TypeError, ValueError):
        total = 0
    if cod and total > 0:
        tem_estoque.add(str(cod).strip().lstrip("0"))
print(f"  estoque: {os.path.basename(arq_est)} -> {len(tem_estoque)} com saldo")

# ---------- 3. cadastro (preco + grupo) ----------
arq_cad = mais_recente(CADASTRO_DIR, "CADASTRO DE PRODUTOS ATUALIZADO*.xlsx")
wb = openpyxl.load_workbook(arq_cad, data_only=True, read_only=True)
ws = wb["CADASTRO"]
print(f"  cadastro: {os.path.basename(arq_cad)}")

linhas, medicamento, sem_foto, sem_estoque, excluidos = [], 0, 0, 0, []
for r in ws.iter_rows(min_row=2, values_only=True):
    ean, cod, desc, lab, grupo, _cf, _cm, _uc, pmc, _fr, _m1, _m2, _pr, _dm, lsit = r[:15]
    if not cod:
        continue
    grupo = (grupo or "").strip()
    if not grupo.startswith("HPC"):      # <-- trava de medicamento
        medicamento += 1
        continue
    if str(lsit).strip().lower() != "true":
        continue
    try:
        preco = float(pmc or 0)
    except (TypeError, ValueError):
        preco = 0
    if preco <= 0:
        continue

    k = str(cod).strip()
    if k.lstrip("0") not in tem_estoque:
        sem_estoque += 1
        continue
    if k not in fotos:
        sem_foto += 1
        continue

    nome_ib, marca, img = fotos[k]
    titulo = (nome_ib or desc or "").strip()
    if any(x in norm(titulo) for x in EXCLUIR_NOME):
        excluidos.append(titulo)
        continue

    linhas.append({
        "id": k,
        "title": titulo[:200],
        "description": titulo[:9999],
        "availability": "in stock",
        "condition": "new",
        "price": f"{preco:.2f} BRL",
        "link": LINK_LOJA,
        "image_link": img,
        "brand": (marca or "").strip(),
    })

# ---------- 4. gravar ----------
# Ordem alfabetica (por titulo) so afeta a ORDEM DE EXIBICAO no catalogo do WhatsApp/loja;
# quais produtos entram continua decidido pelos filtros acima (e, no top 490, pelo faturamento).
cols = ["id", "title", "description", "availability", "condition",
        "price", "link", "image_link", "brand"]
linhas.sort(key=lambda l: norm(l["title"]))
with open(SAIDA, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=cols)
    w.writeheader()
    w.writerows(linhas)

# ---------- 5. top 490 mais vendidos (limite de 500 produtos do WhatsApp Business app) ----------
# Ranking por faturamento (quantidade do fechamento mensal mais recente x P.M.C. do cadastro). Ficam de fora o CD (filial 000) e as
# linhas de movimentacao de estoque (devolucao / remanejo), que nao sao venda ao cliente.
VENDAS_DIR = os.path.join(BASE, "ESTOQUE - VENDAS LOJAS", "VENDAS", "VENDAS GERAL")
SAIDA_TOP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "catalogo_top500.csv")
LIMITE_WHATSAPP = 490
IGNORAR_VENDEDOR = {"devolucao mercadorias", "remanejos entre lojas"}

arqs_ven = sorted(a for a in glob.glob(os.path.join(VENDAS_DIR, "20??-?? Vendas geral*.xlsx"))
                  if "editada" not in norm(os.path.basename(a)))
if not arqs_ven:
    sys.exit(f"ERRO: nenhum arquivo 'AAAA-MM Vendas geral*.xlsx' em {VENDAS_DIR}")
arq_ven = arqs_ven[-1]
wb = openpyxl.load_workbook(arq_ven, data_only=True, read_only=True)
ws = wb["Plan1"] if "Plan1" in wb.sheetnames else wb[wb.sheetnames[0]]
linhas_ven = ws.iter_rows(values_only=True)
cab = [norm(c or "") for c in next(linhas_ven)]
i_fil, i_prod, i_qtd = cab.index("filial"), cab.index("produto"), cab.index("quantidade")
i_vend = cab.index("nome vendedor") if "nome vendedor" in cab else None

vendido = {}
for r in linhas_ven:
    cod, fil = r[i_prod], str(r[i_fil] or "").strip()
    if not cod or not fil or fil.lstrip("0") == "":
        continue
    if i_vend is not None and norm(r[i_vend] or "") in IGNORAR_VENDEDOR:
        continue
    try:
        q = float(r[i_qtd] or 0)
    except (TypeError, ValueError):
        continue
    k = str(cod).strip().lstrip("0")
    vendido[k] = vendido.get(k, 0) + q

def faturamento(l):
    return vendido.get(l["id"].lstrip("0"), 0) * float(l["price"].split()[0])

top = sorted((l for l in linhas if faturamento(l) > 0), key=faturamento, reverse=True)[:LIMITE_WHATSAPP]
top.sort(key=lambda l: norm(l["title"]))  # exibicao alfabetica; entrada no top 500 já foi decidida acima por faturamento
with open(SAIDA_TOP, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=cols)
    w.writeheader()
    w.writerows(top)

print(f"""
  descartados:
    medicamento (LB/GENER/SIMIL): {medicamento}
    sem estoque em nenhuma loja:  {sem_estoque}
    sem foto publica:             {sem_foto}
    excluidos por nome:           {len(excluidos)} {excluidos if excluidos else ''}

  CATALOGO GERADO: {len(linhas)} produtos
  arquivo: {SAIDA}

  CATALOGO WHATSAPP (mais vendidos): {len(top)} produtos
  ranking de vendas: {os.path.basename(arq_ven)}
  arquivo: {SAIDA_TOP}
""")
