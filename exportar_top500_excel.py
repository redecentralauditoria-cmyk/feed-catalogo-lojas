# Gera o Excel (com filtro) dos 500 do catalogo do WhatsApp: marca propria (checklist) primeiro, depois os mais vendidos
# por faturamento = Quantidade x Valor Bruto (fechamento mais recente). Aba 2 = resumo por Grupo FarmaPRO.
import csv, os, glob, unicodedata, collections, datetime, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
A = r"C:\Users\marcel.pereira\CLAUDE\AUTOMATIZ. ATENDIM.  WHATSAPP  LOJAS"
D = A + r"\ATUALIZAÇÃO DIARIA - CATALOGO PROD"
CHECKLIST_MP = r"C:\Users\marcel.pereira\Desktop\AUDITORIAS\01-CLAUDE\VERIFICAR NAS LOJAS\01- CHECKLIST - ULTIMO\Checklist_MarcaPropria_RedeCentral_2026_25.xlsx"
n = lambda s: unicodedata.normalize('NFKD', s.lower()).encode('ascii', 'ignore').decode()
f = max(glob.glob(D + r"\CADASTRO*.xlsx"), key=os.path.getmtime)
g = {str(r[1]).strip().lstrip('0'): (r[4] or '').strip()
     for r in openpyxl.load_workbook(f, read_only=True, data_only=True)["CADASTRO"].iter_rows(min_row=2, values_only=True) if r[1]}
mp = set()
if os.path.exists(CHECKLIST_MP):
    for w in openpyxl.load_workbook(CHECKLIST_MP, read_only=True, data_only=True).worksheets:
        if 'Instru' in w.title or 'Geral' in w.title or w.title.strip().upper() == 'MEDICAMENTO':
            continue
        for r in w.iter_rows(values_only=True):
            if r and len(r) > 1 and r[1] and str(r[1]).strip().isdigit():
                mp.add(str(r[1]).strip().lstrip('0'))
arqs = glob.glob(A + r"\VENDAS GERAL\20??-?? Vendas geral*.xlsx")
vf = max(arqs, key=lambda a: (os.path.basename(a)[:7], "editada" in n(os.path.basename(a))))
it = openpyxl.load_workbook(vf, read_only=True, data_only=True)["Plan1"].iter_rows(values_only=True)
cab = [str(c or '').lower() for c in next(it)]
iv = [i for i, c in enumerate(cab) if 'vendedor' in c and 'nome' in c][0]
ig, ip, iq, ib = cab.index('filial'), cab.index('produto'), cab.index('quantidade'), cab.index('valor bruto')
fat, qt = {}, {}
for r in it:
    if not r[ip] or not r[ig] or str(r[ig]).strip().lstrip('0') == '': continue
    if n(str(r[iv] or '')) in ('devolucao mercadorias', 'remanejos entre lojas'): continue
    k = str(r[ip]).strip().lstrip('0'); q = float(r[iq] or 0)
    fat[k] = fat.get(k, 0) + q * float(r[ib] or 0); qt[k] = qt.get(k, 0) + q
top = list(csv.DictReader(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'catalogo_top500.csv'), encoding='utf-8-sig')))
rows = []
for r in top:
    k = r['id'].lstrip('0')
    rows.append([r['id'], r['title'], r['brand'], g.get(k, ''), float(r['price'].split()[0]),
                 qt.get(k, 0), round(fat.get(k, 0), 2), 'SIM' if k in mp else ''])
rows.sort(key=lambda x: (x[7] != 'SIM', -x[6]))

wb = openpyxl.Workbook(); s = wb.active; s.title = 'Catálogo 500'
s.append(['Posição', 'Código', 'Produto', 'Marca', 'Grupo FarmaPRO', 'Preço (P.M.C.)', 'Qtde vendida (mês)',
          'Faturamento (Qtde x Valor Bruto)', 'Marca própria'])
for i, r in enumerate(rows, 1): s.append([i] + r)
for c in s[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(wrap_text=True, vertical='center')
for i, wd in enumerate([9, 10, 58, 26, 36, 14, 14, 22, 12], 1): s.column_dimensions[get_column_letter(i)].width = wd
s.row_dimensions[1].height = 45
for r in s.iter_rows(min_row=2): r[5].number_format = 'R$ #,##0.00'; r[7].number_format = 'R$ #,##0.00'; r[6].number_format = '#,##0'
s.freeze_panes = 'C2'; s.auto_filter.ref = s.dimensions

r2 = wb.create_sheet('Resumo por categoria')
r2.append(['Grupo FarmaPRO', 'Marca própria', 'Demais (mais vendidos)', 'Total'])
cnt = collections.defaultdict(lambda: [0, 0])
for x in rows: cnt[x[3] or '(sem grupo)'][0 if x[7] == 'SIM' else 1] += 1
for k, (a, b) in sorted(cnt.items(), key=lambda kv: -sum(kv[1])): r2.append([k, a, b, a + b])
r2.append(['TOTAL', sum(v[0] for v in cnt.values()), sum(v[1] for v in cnt.values()), len(rows)])
for c in r2[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
for c in r2[r2.max_row]: c.font = Font(bold=True)
for i, wd in enumerate([40, 14, 22, 10], 1): r2.column_dimensions[get_column_letter(i)].width = wd
r2.freeze_panes = 'A2'

out = A + r"\catalogo_whatsapp_500_produtos_" + datetime.date.today().strftime('%d-%m-%Y') + ".xlsx"
wb.save(out); print(out, len(rows), 'marca propria:', sum(1 for x in rows if x[7] == 'SIM'), os.path.basename(vf))
for k, (a, b) in sorted(cnt.items(), key=lambda kv: -sum(kv[1])): print(f"  {k}: MP {a} + outros {b} = {a+b}")
