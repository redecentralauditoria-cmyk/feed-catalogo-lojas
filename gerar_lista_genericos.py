"""
Gera a LISTA DE PRECOS DE GENERICOS para carregar em Meta Business Agent > Conhecimento > Arquivos e links,
para a IA oferecer o generico (mesmo principio ativo, dosagem e forma) quando o cliente pedir um medicamento
de referencia ou similar (pedido do Marcel em 07/10/2026 - meta de genericos da equipe).

- Nome = descricao do CADASTRO (traz principio ativo, dosagem, apresentacao e laboratorio).
- Preco = P.M.C. do CADASTRO (cadastro Fidelidade). So itens ativos e com estoque em alguma loja. Sem corte por ranking.
- Grupos em GRUPOS_GENERICO. Psicotropicos/antibioticos so entram se INCLUIR_CONTROLADOS = True (decisao do Marcel).
"""
import openpyxl, glob, os, sys, unicodedata
from datetime import date
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

PASTA = r"C:\Users\marcel.pereira\CLAUDE\AUTOMATIZ. ATENDIM.  WHATSAPP  LOJAS"
DIARIO = os.path.join(PASTA, "ATUALIZAÇÃO DIARIA - CATALOGO PROD")
INCLUIR_CONTROLADOS = False
GRUPOS_GENERICO = {"GENER GERAL", "GENER CONTINUOS", "GENER O.T.C"}
if INCLUIR_CONTROLADOS:
    GRUPOS_GENERICO |= {"GENER CONTINUOS PSICOTROPICOS", "GENER ANTIBIOTICOS"}


def norm(s):
    s = unicodedata.normalize("NFD", str(s))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip().lower()


def mais_recente(padrao):
    achados = glob.glob(os.path.join(DIARIO, padrao))
    if not achados:
        sys.exit(f"ERRO: nenhum arquivo '{padrao}' em {DIARIO}")
    return max(achados, key=os.path.getmtime)


arq_est = mais_recente("Estoque das Filiais em *.XLSX")
tem_estoque = set()
for r in openpyxl.load_workbook(arq_est, read_only=True, data_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
    try:
        if r[0] and float(r[3] or 0) > 0:
            tem_estoque.add(str(r[0]).strip().lstrip("0"))
    except (TypeError, ValueError):
        pass

arq_cad = mais_recente("CADASTRO DE PRODUTOS ATUALIZADO*.xlsx")
itens = []
for r in openpyxl.load_workbook(arq_cad, read_only=True, data_only=True)["CADASTRO"].iter_rows(min_row=2, values_only=True):
    cod, desc, grupo, pmc, lsit = r[1], r[2], (r[4] or "").strip(), r[8], r[14]
    if not cod or grupo not in GRUPOS_GENERICO or str(lsit).strip().lower() != "true":
        continue
    try:
        preco = float(pmc or 0)
    except (TypeError, ValueError):
        continue
    if preco <= 0 or str(cod).strip().lstrip("0") not in tem_estoque:
        continue
    itens.append((" ".join(str(desc or "").split()), preco))
itens.sort(key=lambda x: norm(x[0]))

hoje = date.today()
saida = os.path.join(DIARIO, f"precos_genericos_{hoje:%d-%m-%Y}.pdf")
st = ParagraphStyle("c", fontName="Helvetica", fontSize=8.5, leading=11)
stt = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=13, leading=17, spaceAfter=6)
S = [Paragraph("Rede Central Farma - Lista de preços de medicamentos GENÉRICOS", stt),
     Paragraph(f"<b>Atualizada em {hoje:%d/%m/%Y}.</b> Esta lista SUBSTITUI qualquer lista de genéricos anterior. "
               "Preços do cadastro Fidelidade (já com desconto). Cada linha traz o princípio ativo, a dosagem, a apresentação "
               "e o laboratório. Use só para informar o preço do genérico com o MESMO princípio ativo, a MESMA dosagem e a MESMA "
               "forma do medicamento que o cliente pediu - nunca para orientar uso, dose ou indicação. Não é oferta nem promoção.", st),
     Spacer(1, 8)]
dados = [["Genérico (princípio ativo, dosagem, apresentação, laboratório)", "Preço (R$)"]] + \
        [[Paragraph(d, st), f"{p:.2f}".replace(".", ",")] for d, p in itens]
t = Table(dados, colWidths=[14.5 * cm, 2.5 * cm], repeatRows=1)
t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                       ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                       ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5D8DD")), ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                       ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F3F5")])]))
S.append(t)
SimpleDocTemplate(saida, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                  title="Precos genericos - Rede Central Farma").build(S)
print(f"cadastro: {os.path.basename(arq_cad)} | estoque: {os.path.basename(arq_est)}")
print(f"genericos na lista: {len(itens)} | grupos: {sorted(GRUPOS_GENERICO)}")
print(f"arquivo: {saida} ({os.path.getsize(saida)//1024} KB)")
