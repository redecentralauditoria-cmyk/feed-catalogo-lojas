"""
Lista de preco dos DEMAIS produtos (pedido do Marcel 10/10/2026): tudo o que esta ativo no cadastro, com preco e com saldo
em alguma loja, e que NAO esta no catalogo do WhatsApp (catalogo_top500.csv), nem na lista de venda livre, nem na de genericos.
Fora (por enquanto): controlados/psicotropicos (grupos PSICOTROPICOS ou tag (A1..C5,B1,B2) na descricao) e antibioticos
(grupos ANTIBIOTICOS ou tag (AB)); itens de consumo interno (CONSUMO ...), manipulados e vacinas.
Gera 2 PDFs (medicamentos / demais produtos) para Arquivos e links da I.A. Rodar DEPOIS de gerar_catalogo.py e gerar_lista_medicamentos.py.
"""
import openpyxl, glob, os, re, csv, sys, unicodedata
from datetime import date
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

def preco_venda(pmc, promo):
    """Preço real de venda (Marcel 21/10/2026): P.M.C. x (1 - coluna Promoção do cadastro). Promoção é a fração de desconto (0,59 = 59%)."""
    p = float(pmc or 0)
    try: d = float(promo or 0)
    except (TypeError, ValueError): d = 0
    if 0 < d < 1: p = p * (1 - d)
    return round(p, 2)

AQUI = os.path.dirname(os.path.abspath(__file__))
PASTA = r"C:\Users\marcel.pereira\CLAUDE\AUTOMATIZ. ATENDIM.  WHATSAPP  LOJAS"
DIARIO = os.path.join(PASTA, "ATUALIZAÇÃO DIARIA - CATALOGO PROD")
ESTOQUE_MINIMO = 1    # Marcel 21/10: preço de TODOS os itens com estoque na rede (antes 10 un, decisão de 10/10)
GENERICOS = {"GENER GERAL", "GENER CONTINUOS", "GENER O.T.C"}          # ja estao na lista de genericos
TAG_CONTROLADO = re.compile(r"\((A1|A2|A3|B1|B2|C1|C2|C3|C4|C5|AB)\)")
norm = lambda s: "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn").lower().strip()

def recente(p):
    a = glob.glob(os.path.join(DIARIO, p))
    if not a: sys.exit(f"ERRO: falta {p}")
    return max(a, key=os.path.getmtime)

arq_est, arq_cad = recente("Estoque das Filiais em *.XLSX"), recente("CADASTRO DE PRODUTOS ATUALIZADO*.xlsx")
tem = set()
for r in openpyxl.load_workbook(arq_est, read_only=True, data_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
    try:
        if r[0] and float(r[3] or 0) >= ESTOQUE_MINIMO: tem.add(str(r[0]).strip().lstrip("0"))
    except (TypeError, ValueError): pass
ja = {r["id"].lstrip("0") for r in csv.DictReader(open(os.path.join(AQUI, "catalogo_top500.csv"), encoding="utf-8-sig"))}
vl = os.path.join(DIARIO, "codigos_venda_livre.txt")
if not os.path.exists(vl): sys.exit("ERRO: rode antes gerar_lista_medicamentos.py")
ja |= set(open(vl).read().split())
nomes = {}
for r in openpyxl.load_workbook(os.path.join(PASTA, "BANCO_IMAGENS_INSTABUY.xlsx"), read_only=True, data_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
    if r[1] and r[4]: nomes[str(r[1]).strip().lstrip("0")] = str(r[4]).strip().lstrip("*")

med, outros, fora = [], [], {"controlado/antibiotico": 0, "consumo/manipulado/vacina": 0}
for r in openpyxl.load_workbook(arq_cad, read_only=True, data_only=True)["CADASTRO"].iter_rows(min_row=2, values_only=True):
    cod, desc, grupo, pmc, lsit, promo = r[1], str(r[2] or "").strip(), (r[4] or "").strip(), r[8], r[14], r[12]
    if not cod or str(lsit).strip().lower() != "true": continue
    k = str(cod).strip().lstrip("0")
    if k in ja or k not in tem or grupo in GENERICOS: continue
    try: preco = preco_venda(pmc, promo)
    except (TypeError, ValueError): continue
    if preco <= 0: continue
    if "PSICO" in grupo or "ANTIBIOT" in grupo or TAG_CONTROLADO.search(desc.upper()):
        fora["controlado/antibiotico"] += 1; continue
    if grupo.startswith("CONSUMO") or grupo in ("MANIPULADOS", "LB VACINAS"):
        fora["consumo/manipulado/vacina"] += 1; continue
    eh_med = grupo.startswith(("LB ", "SIMIL", "GENER")) and not grupo.startswith(("LB LIBERADOS", "LB EXCESSO ESTOQUE HPC"))
    nome = re.sub(r"\s+", " ", desc) if eh_med else (nomes.get(k) or re.sub(r"\s+", " ", desc).title())
    (med if eh_med else outros).append((nome, preco))

hoje = date.today()
st = ParagraphStyle("c", fontName="Helvetica", fontSize=8, leading=10)
stt = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=13, leading=17, spaceAfter=6)
def pdf(itens, arquivo, titulo, aviso):
    itens.sort(key=lambda x: norm(x[0]))
    S = [Paragraph(titulo, stt), Paragraph(f"<b>Preços de {hoje:%d/%m/%Y}.</b> " + aviso, st), Spacer(1, 6)]
    for n, p in itens:
        S.append(Paragraph(f"{n.replace('&', '&amp;')} - R$ {p:.2f}".replace(".", ",", 0) .replace(f"{p:.2f}", f"{p:.2f}".replace(".", ",")), st))
    saida = os.path.join(DIARIO, arquivo)
    SimpleDocTemplate(saida, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm, title=titulo).build(S)
    print(f"{arquivo}: {len(itens)} itens, {os.path.getsize(saida)//1024} KB")

base = ("Preços do cadastro Fidelidade (já com desconto). Esta lista SUBSTITUI a do dia anterior. Use só para informar preço quando o "
        "cliente pedir o produto pelo nome (mesmo nome, dosagem, forma e tamanho) - nunca para orientar uso, dose ou indicação. ")
pdf(med, f"precos_demais_medicamentos_{hoje:%d-%m-%Y}.pdf", "Rede Central Farma - Preços dos demais medicamentos (referência e similares)",
    base + "Não inclui controlados nem antibióticos. Quando existir genérico do mesmo princípio ativo, dosagem e forma na lista de genéricos, informe também o genérico.")
pdf(outros, f"precos_demais_produtos_{hoje:%d-%m-%Y}.pdf", "Rede Central Farma - Preços dos demais produtos (fora do catálogo)", base)
print(f"cadastro {os.path.basename(arq_cad)} | estoque {os.path.basename(arq_est)} | ja no catalogo/listas: {len(ja)} | fora: {fora}")
