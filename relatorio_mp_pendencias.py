# Relatorio das pendencias da marca propria no catalogo WhatsApp (regra MP em discussao com o Marcel - 07-10/10/2026)
import openpyxl, glob, os, re, unicodedata, datetime
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
A = r"C:\Users\marcel.pereira\CLAUDE\AUTOMATIZ. ATENDIM.  WHATSAPP  LOJAS"
D = A + r"\ATUALIZAÇÃO DIARIA - CATALOGO PROD"
CK = r"C:\Users\marcel.pereira\Desktop\AUDITORIAS\01-CLAUDE\VERIFICAR NAS LOJAS\01- CHECKLIST - ULTIMO\Checklist_MarcaPropria_RedeCentral_2026_25.xlsx"
FOTO_ERRADA = {"115550", "116282"}   # foto/titulo do banco Instabuy e de OUTRO produto (Eudora / Integralmedica)
n = lambda s: unicodedata.normalize('NFKD', str(s).lower()).encode('ascii', 'ignore').decode()
cadf = max(glob.glob(D + r"\CADASTRO*.xlsx"), key=os.path.getmtime)
estf = max(glob.glob(D + r"\Estoque das Filiais em *.XLSX"), key=os.path.getmtime)
cad = {}
for r in openpyxl.load_workbook(cadf, read_only=True, data_only=True)["CADASTRO"].iter_rows(min_row=2, values_only=True):
    if r[1]: cad[str(r[1]).strip().lstrip('0')] = (str(r[2]).strip(), (r[4] or '').strip(), r[8])
est = {}
for r in openpyxl.load_workbook(estf, read_only=True, data_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
    if r[0]:
        try: est[str(r[0]).strip().lstrip('0')] = float(r[3] or 0)
        except: pass
foto = {}
for r in openpyxl.load_workbook(A + r"\BANCO_IMAGENS_INSTABUY.xlsx", read_only=True, data_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
    if r[1]: foto[str(r[1]).strip().lstrip('0')] = (r[4], r[7])
arqs = glob.glob(A + r"\VENDAS GERAL\20??-?? Vendas geral*.xlsx")
vf = max(arqs, key=lambda a: (os.path.basename(a)[:7], "editada" in n(os.path.basename(a))))
it = openpyxl.load_workbook(vf, read_only=True, data_only=True)["Plan1"].iter_rows(values_only=True)
cab = [n(c or '') for c in next(it)]
iv = [i for i, c in enumerate(cab) if 'vendedor' in c and 'nome' in c][0]
ig, ip, iq = cab.index('filial'), cab.index('produto'), cab.index('quantidade')
qt = {}
for r in it:
    if not r[ip] or not r[ig] or str(r[ig]).strip().lstrip('0') == '': continue
    if n(str(r[iv] or '')) in ('devolucao mercadorias', 'remanejos entre lojas'): continue
    k = str(r[ip]).strip().lstrip('0'); qt[k] = qt.get(k, 0) + float(r[iq] or 0)

items = []
for w in openpyxl.load_workbook(CK, read_only=True, data_only=True).worksheets:
    if 'Instru' in w.title or 'Geral' in w.title: continue
    for r in w.iter_rows(values_only=True):
        if r and len(r) > 3 and r[1] and str(r[1]).strip().isdigit():
            items.append((w.title.strip(), str(r[1]).strip().lstrip('0'), str(r[3]).strip()))
rows = []
for cat, k, desc in items:
    c = cad.get(k); e = est.get(k, 0); f = foto.get(k); q = qt.get(k, 0)
    sem_foto = (not f) or (not f[1]) or f[1] == 'Sem imagem'
    natus = 'natus' in n(desc)
    med = cat.upper() == 'MEDICAMENTO'
    rows.append(dict(cat=cat, cod=k, desc=desc, cad=c[0] if c else '', grupo=c[1] if c else '', preco=c[2] if c else None, est=e, q=q,
                     sem_cad=not c, sem_est=bool(c) and e <= 0, sem_foto=sem_foto, foto_errada=k in FOTO_ERRADA, natus=natus, med=med,
                     titulo_banco=f[0] if f else ''))
ok = lambda r: not r['sem_cad'] and not r['sem_est'] and not r['natus'] and not r['med']

wb = openpyxl.Workbook()
def aba(nome, cols, dados, larg):
    s = wb.create_sheet(nome); s.append(cols)
    for d in dados: s.append(d)
    for c in s[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, wd in enumerate(larg, 1): s.column_dimensions[get_column_letter(i)].width = wd
    s.freeze_panes = 'A2'; s.auto_filter.ref = s.dimensions
    return s
wb.remove(wb.active)

# 1 - fotos para o Vinicius
fv = [r for r in rows if ok(r) and (r['sem_foto'] or r['foto_errada'])]
aba('Fotos p Vinicius', ['Código', 'Categoria', 'Produto (checklist)', 'Nome no cadastro', 'Situação da foto', 'Qtde vendida (mês)'],
    [[r['cod'], r['cat'], r['desc'], r['cad'], 'Foto do banco é de OUTRO produto (' + r['titulo_banco'] + ')' if r['foto_errada'] else 'Sem foto', r['q']] for r in sorted(fv, key=lambda r: (r['cat'], r['desc']))],
    [10, 16, 58, 46, 52, 14])
# 2 - menos de 5 un
m5 = [r for r in rows if ok(r) and r['q'] < 5]
aba('Menos de 5 un', ['Código', 'Categoria', 'Produto (checklist)', 'Nome no cadastro', 'Qtde vendida (mês)', 'Estoque (rede)', 'Preço (P.M.C.)', 'Foto'],
    [[r['cod'], r['cat'], r['desc'], r['cad'], r['q'], r['est'], r['preco'], 'sem foto' if r['sem_foto'] else ''] for r in sorted(m5, key=lambda r: (r['q'], r['desc']))],
    [10, 16, 58, 46, 14, 12, 12, 10])
# 3 - fora do cadastro / sem estoque
fo = [r for r in rows if (r['sem_cad'] or r['sem_est']) and not r['natus']]
aba('Sem cadastro ou estoque', ['Código', 'Categoria', 'Produto (checklist)', 'Motivo'],
    [[r['cod'], r['cat'], r['desc'], 'Fora do cadastro' if r['sem_cad'] else 'Sem estoque em nenhuma loja'] for r in fo], [10, 16, 58, 32])
# 4 - medicamento
aba('Medicamento (lista preços)', ['Código', 'Produto (checklist)', 'Nome no cadastro', 'Grupo', 'Preço (P.M.C.)', 'Qtde vendida (mês)', 'Estoque'],
    [[r['cod'], r['desc'], r['cad'], r['grupo'], r['preco'], r['q'], r['est']] for r in rows if r['med']], [10, 50, 44, 22, 12, 14, 10])
# 5 - natusplant
aba('Natusplant (fora)', ['Código', 'Produto (checklist)', 'Qtde vendida (mês)', 'Estoque'],
    [[r['cod'], r['desc'], r['q'], r['est']] for r in rows if r['natus']], [10, 62, 14, 10])
out = A + r"\marca_propria_pendencias_catalogo_" + datetime.date.today().strftime('%d-%m-%Y') + ".xlsx"
wb.save(out)
print(out); print('checklist', len(rows), '| fotos', len(fv), '| <5un', len(m5), '| sem cad/est', len(fo), '| med', sum(r['med'] for r in rows), '| natus', sum(r['natus'] for r in rows), '| vendas', os.path.basename(vf))
for r in fv: print(' FOTO', r['cod'], r['desc'], '| errada' if r['foto_errada'] else '')
