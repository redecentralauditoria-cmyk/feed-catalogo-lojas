"""
Lista de preco de CONTROLADOS (psicotropicos, Portaria 344: tags A1..C5/B1/B2 ou grupo PSICOTROPICOS) e ANTIBIOTICOS
(grupo ANTIBIOTICOS ou tag (AB)) - liberada pelo Marcel em 20/10/2026 para a I.A. ter o preco exato de TODOS os medicamentos.
Inclui referencia, similar e generico (genericos marcados "(genérico)"). Entra se estiver ativo, com preco e com saldo em
alguma loja da rede. Gera precos_controlados_antibioticos_DD-MM-AAAA.pdf para Arquivos e links da I.A.
"""
import openpyxl, glob, os, re, sys, unicodedata
from datetime import date
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

def preco_venda(pmc, promo):
    """Preço real de venda (Marcel 21/10/2026): P.M.C. x (1 - coluna Promoção do cadastro). Promoção é a fração de desconto (0,59 = 59%)."""
    p = float(pmc or 0)
    try: d = float(promo or 0)
    except (TypeError, ValueError): d = 0
    if 0 < d < 1: p = p * (1 - d)
    return round(p, 2)

# Tabloides/ofertas em vigor HOJE (pasta TABELAS DESCONTOS CAMPANHAS) - valem no lugar da coluna Promocao (22/10/2026)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from campanhas import precos_campanha
CAMPANHA = precos_campanha()

def preco_final(cod, pmc, promo):
    p = preco_venda(pmc, promo)
    try: fator = CAMPANHA.get(int(str(cod).strip()))
    except ValueError: fator = None
    c = round(float(pmc or 0) * fator, 2) if fator else 0
    return c if c > 0 and (p <= 0 or c < p) else p

PASTA = r"C:\Users\marcel.pereira\CLAUDE\AUTOMATIZ. ATENDIM.  WHATSAPP  LOJAS"
DIARIO = os.path.join(PASTA, "ATUALIZAÇÃO DIARIA - CATALOGO PROD")
ESTOQUE_MINIMO = 1
TAG_CONTROLADO = re.compile(r"\((A1|A2|A3|B1|B2|C1|C2|C3|C4|C5|AB)\)")
norm = lambda s: "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn").lower().strip()

def recente(p):
    a = glob.glob(os.path.join(DIARIO, p))
    if not a: sys.exit(f"ERRO: falta {p}")
    return max(a, key=os.path.getmtime)

arq_est, arq_cad = recente("Estoque das Filiais em *.XLSX"), recente("CADASTRO DE PRODUTOS ATUALIZADO*.xlsx")
tem = set()
for r in openpyxl.load_workbook(arq_est, read_only=True, data_only=True).worksheets[0].iter_rows(min_row=2, values_only=True):
    try:
        if r[0] and float(r[3] or 0) >= ESTOQUE_MINIMO: tem.add(str(r[0]).strip().lstrip("0"))
    except (TypeError, ValueError): pass

itens, grupos = [], {}
for r in openpyxl.load_workbook(arq_cad, read_only=True, data_only=True)["CADASTRO"].iter_rows(min_row=2, values_only=True):
    cod, desc, grupo, pmc, lsit, promo = r[1], str(r[2] or "").strip(), (r[4] or "").strip(), r[8], r[14], r[12]
    if not cod or str(lsit).strip().lower() != "true": continue
    k = str(cod).strip().lstrip("0")
    if k not in tem: continue
    if not ("PSICO" in grupo or "ANTIBIOT" in grupo or TAG_CONTROLADO.search(desc.upper())): continue
    if grupo.startswith("CONSUMO") or grupo in ("MANIPULADOS", "LB VACINAS"): continue
    try: preco = preco_final(cod, pmc, promo)
    except (TypeError, ValueError): continue
    if preco <= 0: continue
    nome = re.sub(r"\s+", " ", desc)
    if grupo.startswith("GENER") and "GENER" not in nome.upper(): nome += " (genérico)"
    itens.append((nome, preco)); grupos[grupo] = grupos.get(grupo, 0) + 1

hoje = date.today()
st = ParagraphStyle("c", fontName="Helvetica", fontSize=8, leading=10)
stt = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=13, leading=17, spaceAfter=6)
itens.sort(key=lambda x: norm(x[0]))
titulo = "Rede Central Farma - Preços de controlados e antibióticos"
aviso = ("<b>Preços de %s.</b> Preços do cadastro Fidelidade (já com desconto). Esta lista SUBSTITUI a do dia anterior. "
         "Use SOMENTE para informar o preço quando o cliente pedir o medicamento pelo nome, com o MESMO princípio ativo, dosagem, forma e "
         "quantidade. Nunca oriente uso, dose, indicação ou troca, e nunca passe preço de outro princípio ativo ou de outra dosagem. "
         "Controlado: a receita física original é obrigatória (siga a regra de controlado). Itens marcados (genérico) são genéricos." % f"{hoje:%d/%m/%Y}")
S = [Paragraph(titulo, stt), Paragraph(aviso, st), Spacer(1, 6)]
for n, p in itens:
    S.append(Paragraph(f"{n.replace('&', '&amp;')} - R$ {p:.2f}".replace(f"{p:.2f}", f"{p:.2f}".replace(".", ",")), st))
arquivo = f"precos_controlados_antibioticos_{hoje:%d-%m-%Y}.pdf"
saida = os.path.join(DIARIO, arquivo)
SimpleDocTemplate(saida, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm, title=titulo).build(S)
print(f"{arquivo}: {len(itens)} itens, {os.path.getsize(saida)//1024} KB | cadastro {os.path.basename(arq_cad)} | estoque {os.path.basename(arq_est)}")
print("grupos:", dict(sorted(grupos.items(), key=lambda x: -x[1])[:12]))
